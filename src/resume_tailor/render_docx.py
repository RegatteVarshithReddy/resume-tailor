"""Render a TailoredResume and a cover letter to .docx.

Section order, visibility, headings and per-section formatting come from
`tr.layout` (a ResumeLayout). A profile with the default layout renders the
classic five-section resume. The same layout drives render_pdf's fpdf2 fallback,
and LibreOffice converts this DOCX directly to the primary PDF.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.section import WD_SECTION  # noqa: F401  (kept for compat)
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches

from .models import ResumeLayout, SectionSpec, TailoredResume

BODY_FONT = "Calibri"
BASE_SIZE = 10.5


def _hex_to_rgb(h: str) -> RGBColor:
    h = h.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def fmt_period(start: str, end: str, date_format: str = "%b %Y") -> str:
    def one(v: str) -> str:
        v = (v or "").strip()
        if v.lower() in ("present", "current", ""):
            return "Present" if v else ""
        for f in ("%Y-%m", "%Y/%m", "%m/%Y", "%b %Y", "%B %Y", "%Y"):
            try:
                return datetime.strptime(v, f).strftime(date_format or "%b %Y")
            except ValueError:
                continue
        return v

    a, b = one(start), one(end) or "Present"
    return f"{a} – {b}" if a else b


def _base_style(doc: Document) -> None:
    st = doc.styles["Normal"]
    st.font.name = BODY_FONT
    st.font.size = Pt(BASE_SIZE)
    rpr = st.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), BODY_FONT)
    rfonts.set(qn("w:hAnsi"), BODY_FONT)
    for section in doc.sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)


def _content_width(doc: Document):
    s = doc.sections[0]
    return s.page_width - s.left_margin - s.right_margin


def _p(doc, text="", *, size=BASE_SIZE, bold=False, italic=False, color=None,
       align=None, space_before=0, space_after=2, all_caps=False):
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(space_before)
    para.paragraph_format.space_after = Pt(space_after)
    if align is not None:
        para.alignment = align
    if text:
        run = para.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.size = Pt(size)
        if color is not None:
            run.font.color.rgb = color
        if all_caps:
            run.font.all_caps = True
    return para


def _section_heading(doc, title, accent):
    """`accent` is an (r, g, b) tuple."""
    para = _p(doc, title, size=BASE_SIZE + 1.5, bold=True, color=RGBColor(*accent),
              space_before=8, space_after=2, all_caps=True)
    pbdr = para.paragraph_format.element.get_or_add_pPr()
    bdr = pbdr.makeelement(qn("w:pBdr"), {})
    bottom = bdr.makeelement(qn("w:bottom"), {
        qn("w:val"): "single", qn("w:sz"): "6", qn("w:space"): "1",
        qn("w:color"): "%02X%02X%02X" % (accent[0], accent[1], accent[2]),
    })
    bdr.append(bottom)
    pbdr.append(bdr)
    return para


def _bullet(doc, text):
    para = doc.add_paragraph(style=None)
    para.paragraph_format.left_indent = Inches(0.18)
    para.paragraph_format.space_after = Pt(1.5)
    para.add_run("•  ").bold = False
    para.add_run(text)
    return para


def _cap(items, spec: SectionSpec):
    n = spec.max_items
    return items[:n] if n and n > 0 else items


# --------------------------------------------------------------------------- #
# Section renderers — one per core key, plus custom
# --------------------------------------------------------------------------- #
def _sec_summary(doc, tr: TailoredResume, spec: SectionSpec, accent_t) -> None:
    if not tr.summary:
        return
    _section_heading(doc, spec.heading(), accent_t)
    _p(doc, tr.summary, space_after=4)


def _sec_skills(doc, tr: TailoredResume, spec: SectionSpec, accent_t) -> None:
    groups = [g for g in tr.skill_groups if g.skills]
    if not groups:
        return
    _section_heading(doc, spec.heading(), accent_t)
    for g in groups:
        para = doc.add_paragraph()
        para.paragraph_format.space_after = Pt(1.5)
        para.add_run(f"{g.category}: ").bold = True
        para.add_run(", ".join(_cap(g.skills, spec)))


def _sec_experience(doc, tr: TailoredResume, spec: SectionSpec, accent_t) -> None:
    if not tr.experience:
        return
    _section_heading(doc, spec.heading(), accent_t)
    tab_x = _content_width(doc) - Pt(2)
    dfmt = spec.effective_date_format()
    for e in tr.experience:
        header = doc.add_paragraph()
        header.paragraph_format.space_before = Pt(6)
        header.paragraph_format.space_after = Pt(0)
        tabs = header.paragraph_format.tab_stops
        tabs.add_tab_stop(tab_x, alignment=WD_TAB_ALIGNMENT.RIGHT)
        left = e.client + (f"  —  {e.employer}" if e.employer else "")
        rr = header.add_run(left)
        rr.bold = True
        rr.font.size = Pt(BASE_SIZE + 0.5)
        header.add_run("\t" + fmt_period(e.start, e.end, dfmt)).font.size = Pt(BASE_SIZE - 0.5)

        sub = doc.add_paragraph()
        sub.paragraph_format.space_after = Pt(2)
        s2 = sub.add_run(e.role + (f"   |   {e.location}" if e.location else ""))
        s2.italic = True
        s2.font.size = Pt(BASE_SIZE)

        if e.summary:
            _p(doc, e.summary, size=BASE_SIZE - 0.5, italic=True, space_after=2)
        for b in _cap(e.bullets, spec):
            _bullet(doc, b)


def _sec_education(doc, tr: TailoredResume, spec: SectionSpec, accent_t) -> None:
    if not tr.education:
        return
    _section_heading(doc, spec.heading(), accent_t)
    for ed in _cap(tr.education, spec):
        line = ", ".join(x for x in [
            " ".join(x for x in [ed.degree, ed.field_of_study] if x),
            ed.institution, ed.location, ed.year,
        ] if x)
        _p(doc, line, space_after=1.5)


def _sec_certifications(doc, tr: TailoredResume, spec: SectionSpec, accent_t) -> None:
    if not tr.certifications:
        return
    _section_heading(doc, spec.heading(), accent_t)
    for ct in _cap(tr.certifications, spec):
        _bullet(doc, " — ".join(x for x in [ct.name, ct.issuer, ct.year] if x))


def _sec_custom(doc, spec: SectionSpec, accent_t) -> None:
    lines = _cap(spec.content, spec)
    if not lines:
        return
    _section_heading(doc, spec.heading(), accent_t)
    if spec.effective_style() == "paragraph":
        for ln in lines:
            _p(doc, ln, space_after=4)
    else:
        for ln in lines:
            _bullet(doc, ln)


_CORE_RENDERERS = {
    "summary": _sec_summary,
    "skills": _sec_skills,
    "experience": _sec_experience,
    "education": _sec_education,
    "certifications": _sec_certifications,
}


def build_resume_doc(tr: TailoredResume, accent_hex: str = "#1F4E79") -> Document:
    accent = _hex_to_rgb(accent_hex)
    accent_t = (accent[0], accent[1], accent[2])
    doc = Document()
    _base_style(doc)

    c = tr.contact
    _p(doc, c.name, size=20, bold=True, color=accent, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    title = tr.target_role or c.title
    if title:
        _p(doc, title, size=BASE_SIZE + 1, color=accent, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    contact_bits = [x for x in [c.email, c.phone, c.location, c.linkedin, c.website] if x]
    if contact_bits:
        _p(doc, "  |  ".join(contact_bits), size=BASE_SIZE - 1,
           align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    if c.work_authorization:
        _p(doc, f"Work Authorization: {c.work_authorization}", size=BASE_SIZE - 1,
           align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)

    layout = tr.layout or ResumeLayout()
    for spec in layout.visible():
        if spec.is_custom:
            _sec_custom(doc, spec, accent_t)
        else:
            fn = _CORE_RENDERERS.get(spec.key)
            if fn:
                fn(doc, tr, spec, accent_t)

    return doc


def build_cover_letter_doc(tr: TailoredResume, body: str, accent_hex: str = "#1F4E79") -> Document:
    accent = _hex_to_rgb(accent_hex)
    doc = Document()
    _base_style(doc)
    c = tr.contact
    _p(doc, c.name, size=16, bold=True, color=accent, space_after=0)
    bits = [x for x in [c.email, c.phone, c.location, c.linkedin] if x]
    if bits:
        _p(doc, "  |  ".join(bits), size=BASE_SIZE - 1, space_after=6)
    _p(doc, datetime.now().strftime("%B %d, %Y"), space_after=6)
    tgt = " / ".join(x for x in [tr.target_role, tr.target_company] if x)
    if tgt:
        _p(doc, f"Re: {tgt}", bold=True, space_after=6)
    for para in [p for p in body.split("\n\n") if p.strip()]:
        _p(doc, para.strip(), space_after=6)
    return doc


def save_doc(doc: Document, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path
